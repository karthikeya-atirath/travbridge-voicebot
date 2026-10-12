"""
Tool load tester — fires N parallel calls at one (or every) LiveKit function tool
and writes inputs, outputs, status and latency for each call to an Excel file.

The tools are called directly (no LLM, no LiveKit room), exactly as the agent's
AgentSession would execute them, so this does not touch the running agent.

Usage (run from the livekit/ folder):
    python tool_load_tester.py --list                       # show testable tools
    python tool_load_tester.py get_destination_info         # 25 parallel calls
    python tool_load_tester.py get_fare_calendar -n 30      # 30 calls
    python tool_load_tester.py get_travel_package get_package_pricing
    python tool_load_tester.py all -n 20                    # every tool
    python tool_load_tester.py search_packages_by_name --inputs my_inputs.json

Options:
    -n / --calls         number of calls per tool (default 25)
    -c / --concurrency   max calls in flight at once (default = all at once)
    --timeout            per-call timeout in seconds (default 180)
    --inputs FILE        JSON list of kwargs dicts to use instead of the built-in
                         samples (cycled if shorter than -n)
    --use-cache          allow google_tools' in-memory destination cache
                         (off by default, so every call really hits the backend)
    --verbose            also print the tools' own logs to the console

Results: tool_test_results/tool_test_<tools>_<timestamp>.xlsx
         one sheet per tool + a Summary sheet.
Tool logs go to tool_test_results/tool_tests.log (not app_voice.log).

Note: each tool has its own @retry (3 attempts), so a call's latency includes
any internal retries, and EXCEPTION means all 3 attempts raised.
"""

import argparse
import asyncio
import contextvars
import importlib
import itertools
import json
import logging
import os
import statistics
import sys
import time
import traceback
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from typing import Any, Callable, Dict, List, Optional

HERE = os.path.dirname(os.path.abspath(__file__))
RESULTS_DIR = os.path.join(HERE, "tool_test_results")
os.makedirs(RESULTS_DIR, exist_ok=True)

# ── Environment: same .env file the agent loads, but a separate log file ─────
from dotenv import load_dotenv

APP_ENV = os.environ.get("APP_ENV", "prod")  # same default as configuration/config_env.py
load_dotenv(os.path.join(HERE, f".env_{APP_ENV}"), override=True)
os.environ.setdefault("APP_LOG_FILE", os.path.join(RESULTS_DIR, "tool_tests.log"))

if HERE not in sys.path:
    sys.path.insert(0, HERE)

# Stub chat_history: the tools only read chat_history.current_chat_history (None outside
# a live session). The real module needs Elasticsearch/DynamoDB env vars at import time
# and would let test calls write into conversation storage.
import types as _types

_chat_history_stub = _types.ModuleType("chat_history")
_chat_history_stub.current_chat_history = None
sys.modules["chat_history"] = _chat_history_stub

try:
    sys.stdout.reconfigure(encoding="utf-8")  # tool logs contain emoji
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EXCEL_CELL_LIMIT = 32_000  # Excel hard limit is 32,767 chars per cell


# =============================================================================
# SAMPLE INPUTS
# =============================================================================

def _future(days: int) -> str:
    return (datetime.now() + timedelta(days=days)).strftime("%d-%m-%Y")


DESTINATIONS = [
    "Bali", "Dubai", "Singapore", "Thailand", "Paris", "Switzerland", "Maldives",
    "Japan", "Goa", "Kerala", "Vietnam", "London", "Kashmir", "Sri Lanka", "Turkey",
]
HUBS = ["Mumbai", "Delhi", "Hyderabad", "Bengaluru", "Chennai", "Kolkata"]
MONTHS = ["November", "December", "January", "February", "March", "April"]
PACKAGE_IDS = ["PKG014256", "PKG013979", "PKG014042"]
PACKAGE_NAMES = [
    "Singapore", "Bali", "Buy 1 Get 1 Free", "Dubai", "Europe", "Thailand",
    "Fusion Europe", "Bali Split Stay", "Kashmir", "Maldives",
]
PREFERENCES = [
    "", "family friendly", "budget", "luxury", "vegetarian options",
    "adventure", "romantic honeymoon", "with kids", "nightlife",
]
RECOMMEND_QUERIES = [
    "beach holiday in December under 1 lakh",
    "honeymoon destination with mountains",
    "family trip with kids in summer",
    "adventure trip for friends in winter",
    "visa free international trip from India",
    "short 4 day weekend getaway from Mumbai",
    "snow destination in January",
    "cultural trip with history and food",
]
ROOM_CONFIGS = [
    [{"roomNo": 1, "noAdult": 2, "noCwb": 0, "noCnbS": 0, "inf": 0, "pax": 2}],
    [{"roomNo": 1, "noAdult": 2, "noCwb": 0, "noCnbS": 1, "inf": 0, "pax": 3}],
    [{"roomNo": 1, "noAdult": 2, "noCwb": 1, "noCnbS": 0, "inf": 0, "pax": 3}],
    [{"roomNo": 1, "noAdult": 2, "noCwb": 0, "noCnbS": 0, "inf": 0, "pax": 2},
     {"roomNo": 2, "noAdult": 2, "noCwb": 0, "noCnbS": 0, "inf": 0, "pax": 2}],
]
FLIGHT_ROUTES = [
    ("dom", "Hyderabad", "Goa"), ("dom", "Mumbai", "Delhi"), ("dom", "Bengaluru", "Kolkata"),
    ("intl", "Mumbai", "Dubai"), ("intl", "Delhi", "Singapore"), ("intl", "Hyderabad", "Bangkok"),
    ("intl", "Chennai", "Kuala Lumpur"), ("dom", "Delhi", "Srinagar"),
]
ITINERARY_CHANGES = [
    "change Day 2 hotel to 5 star",
    "add scuba diving on Day 3",
    "remove the first sightseeing on Day 1",
    "make the trip more budget friendly",
    "add a sunset cruise on the last day",
    "swap Day 1 and Day 2 activities",
]


def _pick(pool: list, i: int):
    return pool[i % len(pool)]


def gen_get_travel_package(i):
    return {
        "destination": _pick(DESTINATIONS, i),
        "number_of_people": _pick([2, 4, 1, 3], i),
        "days": _pick([5, 7, 4, 10], i),
        "budget": _pick([100000, 200000, 75000, 300000], i),
        "hub": _pick(HUBS, i),
        "month_of_travel": _pick(MONTHS, i),
        "package_type": _pick(["GIT,FIT", "GIT", "FIT"], i),
    }


def gen_get_fare_calendar(i):
    return {"package_id": _pick(PACKAGE_IDS, i), "departure_city": _pick(HUBS + ["Joining Direct"], i)}


def gen_get_package_pricing(i):
    return {
        "pkg_id": _pick(PACKAGE_IDS, i),
        "departure_date": _future(30 + 7 * (i % 10)),
        "hub_city": _pick(HUBS, i),
        "rooms": _pick(ROOM_CONFIGS, i),
        "is_flight_enabled": i % 2 == 1,
    }


def gen_search_packages_by_name(i):
    return {"package_name": _pick(PACKAGE_NAMES, i)}


def gen_destination(i):
    return {"destination": _pick(DESTINATIONS, i)}


def gen_destination_with_pref(i):
    return {"destination": _pick(DESTINATIONS, i), "custom_user_request_input_as_string": _pick(PREFERENCES, i)}


def gen_destination_weather(i):
    return {"destination": _pick(DESTINATIONS, i), "month": _pick(MONTHS, i)}


def gen_budget_guide(i):
    return {"destination": _pick(DESTINATIONS, i), "trip_style": _pick(["budget", "mid-range", "luxury"], i)}


def gen_recommend_destinations(i):
    return {"custom_user_request_input_as_string": _pick(RECOMMEND_QUERIES, i)}


def gen_destination_flights(i):
    trip, src, dst = _pick(FLIGHT_ROUTES, i)
    return {
        "trip": trip, "adult": _pick([1, 2, 2, 3], i), "child": _pick([0, 0, 1], i), "infant": 0,
        "from_city": src, "to_city": dst, "depart": _future(20 + 5 * (i % 8)),
        "custom_input": _pick(["", "cheapest", "non-stop", "early morning"], i),
    }


def gen_create_custom_itinerary(i):
    start = 30 + 3 * (i % 10)
    trip, src, dst = _pick(FLIGHT_ROUTES, i)
    return {
        "destination": dst, "from_city": src,
        "start_date": _future(start), "end_date": _future(start + _pick([3, 4, 5], i)),
        "adults": _pick([2, 2, 1, 4], i), "children": _pick([0, 1, 0], i), "infants": 0,
        "custom_request": _pick(PREFERENCES, i),
    }


# update_custom_itinerary needs a real itinerary → build one first, then reuse it.
async def setup_update_custom_itinerary(runner: "ToolRunner") -> Dict[str, Any]:
    print("  [setup] creating one itinerary via create_custom_itinerary to use as the base...")
    create_tool = _load_tool("google_tools", "create_custom_itinerary")
    res = await create_tool(None, **gen_create_custom_itinerary(0))
    itinerary = res.get("itinerary_result") if isinstance(res, dict) else None
    if not itinerary:
        raise RuntimeError(f"Could not create base itinerary for update test: {str(res)[:300]}")
    return {"existing_itinerary_json": json.dumps(itinerary, ensure_ascii=False)}


def gen_update_custom_itinerary(i, ctx):
    return {"existing_itinerary_json": ctx["existing_itinerary_json"], "change_request": _pick(ITINERARY_CHANGES, i)}


# =============================================================================
# TOOL REGISTRY
# =============================================================================

@dataclass
class ToolSpec:
    module: str
    gen: Callable[..., Dict[str, Any]]
    in_agent: bool                    # registered in AgentSession(tools=[...])
    setup: Optional[Callable] = None  # async fn returning ctx passed to gen(i, ctx)
    notes: str = ""


TOOLS: Dict[str, ToolSpec] = {
    # tools.py
    "get_travel_package":          ToolSpec("tools", gen_get_travel_package, True),
    "get_all_bogo_packages":       ToolSpec("tools", lambda i: {}, True, notes="no inputs"),
    "get_fare_calendar":           ToolSpec("tools", gen_get_fare_calendar, True),
    "get_package_pricing":         ToolSpec("tools", gen_get_package_pricing, True,
                                            notes="dates are synthetic; use --inputs with real bookable dates"),
    "search_packages_by_name":     ToolSpec("tools", gen_search_packages_by_name, True),
    # google_tools.py
    "get_destination_info":        ToolSpec("google_tools", gen_destination, True),
    "get_destination_food":        ToolSpec("google_tools", gen_destination_with_pref, True),
    "get_destination_weather":     ToolSpec("google_tools", gen_destination_weather, True),
    "get_destination_sightseeing": ToolSpec("google_tools", gen_destination_with_pref, True),
    "get_destination_activities":  ToolSpec("google_tools", gen_destination_with_pref, True),
    "get_destination_visa_info":   ToolSpec("google_tools", gen_destination, True),
    "get_destination_hotels":      ToolSpec("google_tools", gen_destination_with_pref, True),
    "recommend_destinations":      ToolSpec("google_tools", gen_recommend_destinations, True),
    "get_destination_flights":     ToolSpec("google_tools", gen_destination_flights, False),
    "get_destination_budget_guide": ToolSpec("google_tools", gen_budget_guide, False),
    "create_custom_itinerary":     ToolSpec("google_tools", gen_create_custom_itinerary, False, notes="slow"),
    "update_custom_itinerary":     ToolSpec("google_tools", gen_update_custom_itinerary, False,
                                            setup=setup_update_custom_itinerary,
                                            notes="creates 1 base itinerary first"),
}


def _load_tool(module_name: str, tool_name: str):
    return getattr(importlib.import_module(module_name), tool_name)


def print_tool_list():
    print("\nTestable tools:\n")
    print(f"  {'#':>2}  {'tool':<30} {'file':<22} {'in agent':<9} notes")
    print(f"  {'-' * 2}  {'-' * 30} {'-' * 22} {'-' * 9} {'-' * 30}")
    for n, (name, spec) in enumerate(TOOLS.items(), 1):
        print(f"  {n:>2}  {name:<30} {spec.module + '.py':<22} {'yes' if spec.in_agent else 'no':<9} {spec.notes}")
    print("\n  'all' runs every tool.\n")


# =============================================================================
# RUNNER
# =============================================================================

@dataclass
class CallResult:
    call_no: int
    started_at: str
    latency_s: float
    status: str               # SUCCESS | TOOL_ERROR | EXCEPTION | TIMEOUT
    code: Any
    message: str
    inputs: Dict[str, Any]
    output: Any
    error_logs: List[str] = field(default_factory=list)


# Each parallel call runs in its own asyncio task (own contextvars copy), so log
# records emitted inside a tool can be attributed to the call that produced them.
_current_call_logs: contextvars.ContextVar[Optional[List[str]]] = contextvars.ContextVar(
    "_current_call_logs", default=None
)


class _CallLogCapture(logging.Handler):
    def __init__(self):
        super().__init__(level=logging.WARNING)

    def emit(self, record: logging.LogRecord) -> None:
        bucket = _current_call_logs.get()
        if bucket is None:
            return
        line = f"[{record.levelname}] {record.getMessage()}"
        if record.exc_info:
            line += "\n" + "".join(traceback.format_exception(*record.exc_info))
        bucket.append(line)


def classify(output: Any) -> tuple:
    """Return (status, code, message) for a tool's returned value."""
    if not isinstance(output, dict):
        return "SUCCESS", None, ""
    code = output.get("code")
    msg = str(output.get("message") or output.get("error") or "")
    if isinstance(code, int) and code >= 400:
        return "TOOL_ERROR", code, msg
    if output.get("status") == "failed" or "error" in output:
        return "TOOL_ERROR", code, msg
    return "SUCCESS", code, msg


def _unwrap_exception(exc: BaseException) -> BaseException:
    # tenacity.RetryError hides the real exception inside last_attempt
    last = getattr(exc, "last_attempt", None)
    if last is not None:
        try:
            inner = last.exception()
            if inner is not None:
                return inner
        except Exception:
            pass
    return exc


class ToolRunner:
    def __init__(self, calls: int, concurrency: Optional[int], timeout: float,
                 custom_inputs: Optional[List[Dict[str, Any]]]):
        self.calls = calls
        self.concurrency = concurrency or calls
        self.timeout = timeout
        self.custom_inputs = custom_inputs

    async def run_tool(self, name: str) -> tuple:
        spec = TOOLS[name]
        tool = _load_tool(spec.module, name)

        if self.custom_inputs:
            inputs = [dict(x) for x in itertools.islice(itertools.cycle(self.custom_inputs), self.calls)]
        else:
            ctx = await spec.setup(self) if spec.setup else None
            inputs = [spec.gen(i, ctx) if spec.setup else spec.gen(i) for i in range(self.calls)]

        sem = asyncio.Semaphore(self.concurrency)

        async def one(i: int, kwargs: Dict[str, Any]) -> CallResult:
            async with sem:
                logs: List[str] = []
                _current_call_logs.set(logs)
                started = datetime.now().strftime("%H:%M:%S.%f")[:-3]
                t0 = time.perf_counter()
                try:
                    out = await asyncio.wait_for(tool(None, **kwargs), timeout=self.timeout)
                    status, code, msg = classify(out)
                except asyncio.TimeoutError:
                    out, status, code, msg = None, "TIMEOUT", None, f"no response within {self.timeout}s"
                except Exception as e:
                    real = _unwrap_exception(e)
                    out = "".join(traceback.format_exception(type(real), real, real.__traceback__))
                    status, code, msg = "EXCEPTION", None, f"{type(real).__name__}: {real}"
                latency = time.perf_counter() - t0
                print(f"  #{i + 1:<3} {status:<10} {latency:7.2f}s  {msg[:80]}")
                return CallResult(i + 1, started, latency, status, code, msg, kwargs, out, logs)

        print(f"\n▶ {name}: {self.calls} calls, concurrency {self.concurrency}")
        wall0 = time.perf_counter()
        results = await asyncio.gather(*(one(i, kw) for i, kw in enumerate(inputs)))
        wall = time.perf_counter() - wall0
        return list(results), wall


def summarize(name: str, results: List[CallResult], wall: float) -> Dict[str, Any]:
    lat = sorted(r.latency_s for r in results)

    def pct(p):
        if not lat:
            return 0.0
        k = min(len(lat) - 1, max(0, round(p / 100 * (len(lat) - 1))))
        return lat[k]

    counts = {s: sum(r.status == s for r in results) for s in ("SUCCESS", "TOOL_ERROR", "EXCEPTION", "TIMEOUT")}
    ok = [r.latency_s for r in results if r.status == "SUCCESS"]
    return {
        "tool": name,
        "calls": len(results),
        **{s.lower(): c for s, c in counts.items()},
        "success_rate_%": round(100 * counts["SUCCESS"] / len(results), 1) if results else 0,
        "min_s": round(lat[0], 3) if lat else 0,
        "avg_s": round(statistics.mean(lat), 3) if lat else 0,
        "p50_s": round(pct(50), 3),
        "p90_s": round(pct(90), 3),
        "p95_s": round(pct(95), 3),
        "max_s": round(lat[-1], 3) if lat else 0,
        "avg_success_s": round(statistics.mean(ok), 3) if ok else None,
        "wall_clock_s": round(wall, 3),
    }


# =============================================================================
# EXCEL
# =============================================================================

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
STATUS_FILLS = {
    "SUCCESS": PatternFill("solid", fgColor="C6EFCE"),
    "TOOL_ERROR": PatternFill("solid", fgColor="FFEB9C"),
    "EXCEPTION": PatternFill("solid", fgColor="FFC7CE"),
    "TIMEOUT": PatternFill("solid", fgColor="FFC7CE"),
}


def _cell_text(value: Any) -> str:
    text = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, default=str)
    return text if len(text) <= EXCEL_CELL_LIMIT else text[:EXCEL_CELL_LIMIT] + " …[truncated]"


def _write_header(ws, headers, widths):
    ws.append(headers)
    for col, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A2"


def write_excel(path: str, runs: Dict[str, tuple], summaries: List[Dict[str, Any]]):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    keys = list(summaries[0].keys())
    _write_header(ws, keys, [30] + [13] * (len(keys) - 1))
    for s in summaries:
        ws.append([s[k] for k in keys])

    headers = ["call_no", "started_at", "latency_s", "status", "code", "message",
               "input_params", "output", "output_chars", "error_logs"]
    for name, (results, _) in runs.items():
        ws = wb.create_sheet(name[:31])
        _write_header(ws, headers, [8, 13, 10, 12, 7, 45, 60, 90, 12, 90])
        for r in sorted(results, key=lambda x: x.call_no):
            out_text = _cell_text(r.output)
            ws.append([r.call_no, r.started_at, round(r.latency_s, 3), r.status, r.code,
                       r.message, _cell_text(r.inputs), out_text, len(out_text),
                       _cell_text("\n".join(r.error_logs))])
            row = ws.max_row
            ws.cell(row=row, column=4).fill = STATUS_FILLS.get(r.status, PatternFill())
            for col in (6, 7, 8, 10):
                ws.cell(row=row, column=col).alignment = Alignment(wrap_text=False, vertical="top")
        ws.auto_filter.ref = ws.dimensions

    wb.save(path)


# =============================================================================
# MAIN
# =============================================================================

def _configure_tool_logs(verbose: bool):
    """UTF-8 log file, optional console silence, and per-call warning/error capture."""
    capture = _CallLogCapture()
    for lg in [logging.getLogger(n) for n in logging.root.manager.loggerDict] + [logging.root]:
        if not isinstance(lg, logging.Logger):
            continue
        for h in list(lg.handlers):
            if isinstance(h, logging.FileHandler) and (h.encoding or "").lower() != "utf-8":
                h.acquire()
                try:
                    h.close()
                    h.encoding = "utf-8"  # Windows defaults to cp1252 and chokes on emoji
                    h.stream = h._open()
                finally:
                    h.release()
            elif type(h) is logging.StreamHandler and not verbose:
                lg.removeHandler(h)
        if lg.name.startswith("voice_bot") or lg is logging.root:
            lg.addHandler(capture)


async def main_async(args) -> int:
    if args.list or not args.tools:
        print_tool_list()
        return 0

    if args.tools == ["all"]:
        names = list(TOOLS)
    else:
        unknown = [t for t in args.tools if t not in TOOLS]
        if unknown:
            print(f"Unknown tool(s): {', '.join(unknown)}")
            print_tool_list()
            return 2
        names = args.tools

    custom_inputs = None
    if args.inputs:
        with open(args.inputs, encoding="utf-8") as f:
            custom_inputs = json.load(f)
        if not isinstance(custom_inputs, list) or not custom_inputs:
            print("--inputs must be a JSON list of kwargs objects")
            return 2

    # Import tool modules now so their loggers exist, then silence the console output.
    for n in names:
        try:
            importlib.import_module(TOOLS[n].module)
        except Exception as e:
            print(f"Cannot import {TOOLS[n].module} for {n}: {e}")
            return 1
    _configure_tool_logs(args.verbose)

    if not args.use_cache:
        import google_tools
        google_tools._check_dest_cache = lambda *a, **k: None  # force real backend calls

    runner = ToolRunner(args.calls, args.concurrency, args.timeout, custom_inputs)
    runs, summaries = {}, []
    for name in names:
        try:
            results, wall = await runner.run_tool(name)
        except Exception as e:
            print(f"  ✖ {name} could not run: {e}")
            continue
        runs[name] = (results, wall)
        s = summarize(name, results, wall)
        summaries.append(s)
        print(f"  = {s['success']}/{s['calls']} success, {s['tool_error']} tool errors, "
              f"{s['exception']} exceptions, {s['timeout']} timeouts | "
              f"avg {s['avg_s']}s p95 {s['p95_s']}s max {s['max_s']}s | wall {s['wall_clock_s']}s")

    if not runs:
        print("Nothing ran.")
        return 1

    label = "all" if args.tools == ["all"] else "_".join(names)[:80]
    path = os.path.join(RESULTS_DIR, f"tool_test_{label}_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    write_excel(path, runs, summaries)
    print(f"\n✔ Results saved to {path}")
    return 0


def main():
    p = argparse.ArgumentParser(description="Parallel load tester for LiveKit agent tools")
    p.add_argument("tools", nargs="*", help="tool name(s), or 'all'")
    p.add_argument("-n", "--calls", type=int, default=25)
    p.add_argument("-c", "--concurrency", type=int, default=None)
    p.add_argument("--timeout", type=float, default=180.0)
    p.add_argument("--inputs", help="JSON file with a list of kwargs dicts")
    p.add_argument("--use-cache", action="store_true")
    p.add_argument("--verbose", action="store_true")
    p.add_argument("--list", action="store_true")
    args = p.parse_args()
    sys.exit(asyncio.run(main_async(args)))


if __name__ == "__main__":
    main()
